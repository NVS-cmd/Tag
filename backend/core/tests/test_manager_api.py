"""API кабинета управляющего: `make test`."""
import json
import uuid
from datetime import timedelta
from io import BytesIO

from django.core.cache import cache
from django.core.management import call_command
from django.test import Client, TestCase
from django.utils import timezone

from core import services
from core.models import AuditLog, Device, Employee, PaymentMethod, Point, Room, User

RUB = 100
PASSWORD = "demo12345demo"


class ManagerApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", verbosity=0)

    def setUp(self):
        cache.clear()
        self.c = Client()
        self.post("/api/login/", {"username": "manager", "password": PASSWORD})
        self.olga = User.objects.get(username="olga")
        self.manager = User.objects.get(username="manager")

    def post(self, url, data=None, client=None):
        return (client or self.c).post(url, json.dumps(data or {}), content_type="application/json")

    def test_admin_cannot_use_manager_api(self):
        admin = Client()
        self.post("/api/login/", {"username": "olga", "password": PASSWORD}, admin)
        self.assertEqual(admin.get("/api/m/summary/").status_code, 403)
        self.assertEqual(admin.get("/api/m/report/").status_code, 403)

    def test_summary_shows_open_and_stale_shift(self):
        shift = services.open_shift(self.olga, 0)
        r = self.c.get("/api/m/summary/").json()
        p1 = next(p for p in r["points"] if p["name"] == "Точка 1")
        self.assertEqual(p1["status"], "open")
        type(shift).objects.filter(pk=shift.pk).update(opened_at=timezone.now() - timedelta(hours=20))
        r = self.c.get("/api/m/summary/").json()
        self.assertEqual(next(p for p in r["points"] if p["name"] == "Точка 1")["status"], "stale")
        self.assertTrue(any(a["tag"] == "Смена" for a in r["alerts"]))

    def test_point_range_30_60_and_validation(self):
        p = Point.objects.get(name="Точка 1")
        self.assertEqual((float(p.rate_min), float(p.rate_max)), (30.0, 60.0))
        r = self.post(f"/api/m/points/{p.id}/", {"rate_min": 35, "rate_max": 30, "rate_default": 32})
        self.assertEqual(r.status_code, 400)
        r = self.post(f"/api/m/points/{p.id}/", {"rate_min": 30, "rate_max": 65, "rate_default": 50})
        self.assertEqual(r.status_code, 200)
        p.refresh_from_db()
        self.assertEqual(float(p.rate_max), 65.0)

    def test_staff_create_update_archive(self):
        p = Point.objects.get(name="Точка 1")
        r = self.post("/api/m/staff/new/", {"kind": "master", "name": "Новая", "point_id": p.id, "rate": 70,
                                            "rate_with_dispatcher": 40, "base_full": False})
        self.assertEqual(r.status_code, 200, r.content)  # управляющий может вне диапазона
        eid = r.json()["employee"]["id"]
        r = self.post(f"/api/m/staff/{eid}/", {"rate": 55})
        self.assertEqual(r.json()["employee"]["rate"], 55.0)
        self.assertTrue(AuditLog.objects.filter(action__startswith="Изменил карточку").exists())
        r = self.post(f"/api/m/staff/{eid}/archive/")
        self.assertTrue(r.json()["employee"]["archived"])
        r = self.post(f"/api/m/staff/{eid}/archive/", {"restore": True})
        self.assertFalse(r.json()["employee"]["archived"])

    def test_admin_rates_create_password_archive(self):
        r = self.post(f"/api/m/admins/{self.olga.id}/", {"rate_session": 12, "categories": {"bar": 15}})
        self.assertEqual(r.status_code, 200, r.content)
        olga = next(a for a in r.json()["admins"] if a["username"] == "olga")
        self.assertEqual((olga["rate_session"], olga["categories"]["bar"]), (12.0, 15.0))
        p = Point.objects.get(name="Точка 3")
        r = self.post("/api/m/admins/new/", {"username": "Nina", "first_name": "Нина", "point_id": p.id,
                                             "password": "short"})
        self.assertEqual(r.status_code, 400)
        r = self.post("/api/m/admins/new/", {"username": "Nina", "first_name": "Нина", "point_id": p.id,
                                             "password": "Kassa-2026-tochka3"})
        self.assertEqual(r.status_code, 200, r.content)
        nina = User.objects.get(username="nina")
        Device.objects.create(user=nina, label="Android")
        self.post(f"/api/m/admins/{nina.id}/password/", {"password": "Novyi-parol-2026"})
        self.assertFalse(Device.objects.get(user=nina).is_active)
        self.post(f"/api/m/admins/{nina.id}/archive/")
        nina.refresh_from_db()
        self.assertFalse(nina.is_active)
        self.post(f"/api/m/admins/{nina.id}/restore/", {"password": "Snova-v-stroyu-2026"})
        nina.refresh_from_db()
        self.assertTrue(nina.is_active)

    def test_devices_revoke_all_keeps_manager(self):
        mine = Device.objects.filter(user=self.manager).first()
        other = Device.objects.create(user=self.olga, label="iPhone")
        self.assertEqual(self.post("/api/m/devices/revoke-all/", {}).status_code, 400)
        r = self.post("/api/m/devices/revoke-all/", {"confirm": True})
        self.assertEqual(r.status_code, 200)
        other.refresh_from_db()
        self.assertFalse(other.is_active)
        if mine:
            mine.refresh_from_db()
            self.assertTrue(mine.is_active)  # сам управляющий не выбит
        r = self.post(f"/api/m/devices/{other.id}/toggle/")
        other.refresh_from_db()
        self.assertTrue(other.is_active)

    def test_manager_adds_cover_master_to_shift(self):
        shift = services.open_shift(self.olga, 0)
        marina = Employee.objects.get(name="Марина")  # мастер Точки 2
        r = self.post(f"/api/m/shifts/{shift.id}/staff/", {"employee_id": marina.id, "present": True})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertIn(marina.id, r.json()["shift"]["attendance"])

    def test_report_and_excel(self):
        shift = services.open_shift(self.olga, 1000 * RUB)
        alina = Employee.objects.get(name="Алина")
        services.mark_attendance(self.olga, shift, alina)
        t0 = timezone.now().replace(second=0, microsecond=0)
        services.create_session(self.olga, shift, room=Room.objects.get(point=self.olga.point, name="Комната 1"),
                                master=alina, start=t0, end=t0 + timedelta(hours=1), amount=6000 * RUB,
                                payment_method=PaymentMethod.CARD, id=uuid.uuid4())
        services.close_shift(self.olga, shift, paid={("employee", alina.id): 3000 * RUB, ("admin", self.olga.id): 600 * RUB},
                             actual_cash=0, comment="тест")
        today = timezone.localdate().isoformat()
        r = self.c.get(f"/api/m/report/?from={today}&to={today}").json()
        p1 = next(p for p in r["points"] if p["point"] == "Точка 1")
        self.assertEqual((p1["revenue"], p1["staff"], p1["profit"]), (6000 * RUB, 3600 * RUB, 2400 * RUB))
        self.assertTrue(any(s["name"] == "Алина" and s["accrued"] == 3000 * RUB for s in r["staff"]))
        x = self.c.get(f"/api/m/report.xlsx?from={today}&to={today}")
        self.assertEqual(x.status_code, 200)
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Точки", "Сотрудники"])

    def test_log(self):
        services.audit(self.manager, "Проверка журнала")
        r = self.c.get("/api/m/log/?q=Проверка").json()
        self.assertEqual(r["log"][0]["action"], "Проверка журнала")
