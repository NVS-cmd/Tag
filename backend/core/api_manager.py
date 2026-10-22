"""JSON API кабинета управляющего. Все суммы — копейки."""
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db.models import Q, Sum
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST

from . import calc, services
from .api import _uuid, api, shift_json
from .models import (
    AuditLog, Device, Employee, Expense, Payout, Point, ProductCategory, RecordStatus, Sale, Session,
    Shift, User,
)

MANAGER = [User.Role.MANAGER, User.Role.SUPERADMIN]
STALE_HOURS = 18
f = float


def _name(u):
    return (u.first_name or u.username) if u else "—"


def _sum(qs, field="amount"):
    return qs.aggregate(t=Sum(field))["t"] or 0


def _day_bounds(d_from: date, d_to: date):
    tz = timezone.get_current_timezone()
    return (timezone.make_aware(datetime.combine(d_from, time.min), tz),
            timezone.make_aware(datetime.combine(d_to + timedelta(days=1), time.min), tz))


# ---------- сводка ----------

@require_GET
@api(roles=MANAGER)
def summary(request):
    now = timezone.now()
    start, end = _day_bounds(timezone.localdate(), timezone.localdate())
    points = []
    for p in Point.objects.filter(is_active=True):
        open_shift = p.shifts.filter(status=Shift.Status.OPEN).select_related("admin").first()
        last = p.shifts.filter(status=Shift.Status.CLOSED).order_by("-closed_at").first()
        today = Shift.objects.filter(point=p).filter(Q(opened_at__gte=start, opened_at__lt=end) | Q(status=Shift.Status.OPEN))
        sess = Session.objects.filter(shift__in=today, status=RecordStatus.OK)
        sales = Sale.objects.filter(shift__in=today, status=RecordStatus.OK)
        stale = bool(open_shift and (now - open_shift.opened_at).total_seconds() > STALE_HOURS * 3600)
        points.append({
            "id": p.id, "name": p.name,
            "status": "stale" if stale else ("open" if open_shift else "closed"),
            "shift_id": open_shift.id if open_shift else None,
            "opened_at": open_shift.opened_at.isoformat() if open_shift else None,
            "admin": _name(open_shift.admin) if open_shift else (_name(last.admin) if last else "—"),
            "on_shift": open_shift.attendance.count() if open_shift else 0,
            "sessions": sess.count(), "revenue": _sum(sess) + _sum(sales),
            "last_discrepancy": last.discrepancy if (last and not open_shift) else None,
        })
    alerts = []
    for s in Shift.objects.filter(status=Shift.Status.CLOSED, closed_at__gte=now - timedelta(days=2)).exclude(
            discrepancy=0).exclude(discrepancy__isnull=True).select_related("point", "admin")[:20]:
        kind = "недостача" if s.discrepancy < 0 else "излишек"
        amount = f"{abs(s.discrepancy) / 100:,.0f}".replace(",", " ")
        alerts.append({"tag": "Касса", "level": "warn", "shift_id": s.id,
                       "text": f"{s.point.name}: {kind} {amount} ₽ при закрытии ({_name(s.admin)}). "
                               f"Причина: «{s.close_comment}»"})
    for p in points:
        if p["status"] == "stale":
            alerts.append({"tag": "Смена", "level": "bad",
                           "text": f"{p['name']}: смена открыта больше {STALE_HOURS} часов и не закрыта."})
    for d in Device.objects.filter(first_seen__gte=now - timedelta(days=1)).select_related("user"):
        if d.user.role == User.Role.ADMIN:
            alerts.append({"tag": "Вход", "level": "info",
                           "text": f"{_name(d.user)} вошла с нового устройства "
                                   f"{timezone.localtime(d.first_seen):%d.%m %H:%M}."})
    return JsonResponse({"points": points, "alerts": alerts})


# ---------- справочник для форм ----------

def _points():
    return [{"id": p.id, "name": p.name, "rate_min": f(p.rate_min), "rate_max": f(p.rate_max),
             "rate_default": f(p.rate_default)} for p in Point.objects.filter(is_active=True)]


@require_GET
@api(roles=MANAGER)
def points_list(request):
    return JsonResponse({"points": _points()})


@require_POST
@api(roles=MANAGER)
def point_update(request, pk):
    d = request.data
    services.update_point_range(request.user, Point.objects.get(pk=pk), rate_min=d.get("rate_min"),
                                rate_max=d.get("rate_max"), rate_default=d.get("rate_default"))
    return JsonResponse({"points": _points()})


# ---------- сотрудники ----------

def _employee_json(e):
    return {
        "id": e.id, "kind": e.kind, "name": e.name, "phone": e.phone, "point_id": e.point_id,
        "point": e.point.name if e.point_id else "Все точки", "rate": f(e.rate),
        "rate_with_dispatcher": f(e.rate_with_dispatcher), "base_full": e.base_full,
        "archived": e.archived_at is not None, "debt": services.employee_debt(e),
    }


@require_GET
@api(roles=MANAGER)
def staff_list(request):
    emps = Employee.objects.select_related("point").order_by("name")
    return JsonResponse({"staff": [_employee_json(e) for e in emps], "points": _points()})


@require_POST
@api(roles=MANAGER)
def staff_create(request):
    d = request.data
    kind = d.get("kind") if d.get("kind") in Employee.Kind.values else Employee.Kind.MASTER
    point = Point.objects.get(pk=d["point_id"]) if d.get("point_id") else None
    if kind == Employee.Kind.MASTER and point is None:
        raise ValidationError("Мастеру нужна точка.")
    rate = services._pct_value(d.get("rate", 50), "Ставка")
    rwd = services._pct_value(d.get("rate_with_dispatcher", rate), "Ставка при диспетчере")
    emp = services.add_employee(request.user, kind=kind, name=str(d.get("name", "")), phone=str(d.get("phone", ""))[:32],
                                point=point, rate=rate, rate_with_dispatcher=rwd, base_full=bool(d.get("base_full", True)))
    return JsonResponse({"employee": _employee_json(emp)})


@require_POST
@api(roles=MANAGER)
def staff_update(request, pk):
    emp = services.update_employee(request.user, Employee.objects.get(pk=pk), request.data)
    return JsonResponse({"employee": _employee_json(emp)})


@require_POST
@api(roles=MANAGER)
def staff_archive(request, pk):
    emp = Employee.objects.get(pk=pk)
    if request.data.get("restore"):
        services.restore_employee(request.user, emp)
    else:
        services.archive_employee(request.user, emp)
    emp.refresh_from_db()
    return JsonResponse({"employee": _employee_json(emp)})


@require_POST
@api(roles=MANAGER)
def shift_add_staff(request, pk):
    """Управляющий добавляет мастера в открытую смену любой точки (или снимает)."""
    shift = Shift.objects.get(pk=pk)
    emp = Employee.objects.get(pk=request.data.get("employee_id"))
    services.mark_attendance(request.user, shift, emp, present=bool(request.data.get("present", True)))
    return JsonResponse({"shift": shift_json(shift)})


@require_GET
@api(roles=MANAGER)
def shift_view(request, pk):
    shift = Shift.objects.select_related("point", "admin").get(pk=pk)
    data = shift_json(shift)
    data.update({"point": shift.point.name, "admin": _name(shift.admin)})
    masters = Employee.objects.filter(kind=Employee.Kind.MASTER, archived_at__isnull=True).select_related("point")
    return JsonResponse({"shift": data, "masters": [
        {"id": m.id, "name": m.name, "point": m.point.name if m.point_id else "", "home": m.point_id == shift.point_id}
        for m in masters]})


# ---------- администраторы ----------

def _admin_json(u, devices_by_user):
    cats = {c.category: f(c.rate) for c in u.category_rates.all()}
    return {
        "id": u.id, "username": u.username, "name": _name(u), "point_id": u.point_id,
        "point": u.point.name if u.point_id else "", "archived": u.archived_at is not None,
        "rate_session": f(u.rate_session), "rate_session_dispatcher": f(u.rate_session_dispatcher),
        "rate_sauna": f(u.rate_sauna), "categories": {c: cats.get(c, 0.0) for c in ProductCategory.values},
        "shift": (lambda s: {"id": s.id, "stale": (timezone.now() - s.opened_at).total_seconds() > STALE_HOURS * 3600}
                  if s else None)(u.shifts.filter(status=Shift.Status.OPEN).first()),
        "devices": devices_by_user.get(u.id, 0),
    }


def _admins_payload():
    devs = defaultdict(int)
    for d in Device.objects.filter(is_active=True):
        devs[d.user_id] += 1
    admins = User.objects.filter(role=User.Role.ADMIN).select_related("point").prefetch_related("category_rates") \
        .order_by("first_name")
    return {"admins": [_admin_json(u, devs) for u in admins], "points": _points(),
            "categories": [{"id": c, "name": n} for c, n in ProductCategory.choices]}


@require_GET
@api(roles=MANAGER)
def admins_list(request):
    return JsonResponse(_admins_payload())


@require_POST
@api(roles=MANAGER)
def admin_create(request):
    d = request.data
    services.create_admin(request.user, username=d.get("username", ""), first_name=d.get("first_name", ""),
                          point=Point.objects.get(pk=d.get("point_id")), password=str(d.get("password", "")))
    return JsonResponse(_admins_payload())


@require_POST
@api(roles=MANAGER)
def admin_update(request, pk):
    services.update_admin(request.user, User.objects.get(pk=pk, role=User.Role.ADMIN), request.data)
    return JsonResponse(_admins_payload())


@require_POST
@api(roles=MANAGER)
def admin_action(request, pk, action):
    admin = User.objects.get(pk=pk, role=User.Role.ADMIN)
    d = request.data
    if action == "archive":
        services.archive_admin(request.user, admin)
    elif action == "restore":
        services.restore_admin(request.user, admin, str(d.get("password", "")))
    elif action == "password":
        services.set_password(request.user, admin, str(d.get("password", "")))
    elif action == "revoke":
        services.revoke_devices(request.user, Device.objects.filter(user=admin), reason="завершены сессии")
    else:
        raise ValidationError("Неизвестное действие.")
    return JsonResponse(_admins_payload())


# ---------- устройства ----------

def _devices_payload(me):
    now = timezone.now()
    rows = []
    for d in Device.objects.select_related("user", "user__point").order_by("-last_seen"):
        rows.append({"id": str(d.id), "who": _name(d.user), "role": d.user.get_role_display(),
                     "point": d.user.point.name if d.user.point_id else "—", "label": _device_label(d.label),
                     "last_seen": d.last_seen.isoformat(), "active": d.is_active, "is_me": d.user_id == me.id,
                     "is_new": (now - d.first_seen).total_seconds() < 86400})
    return {"devices": rows}


def _device_label(ua):
    ua = ua or ""
    for key, name in (("iPhone", "iPhone"), ("iPad", "iPad"), ("Android", "Android"), ("Windows", "Windows"),
                      ("Macintosh", "Mac"), ("Linux", "Linux")):
        if key in ua:
            return name
    return ua[:30] or "Неизвестно"


@require_GET
@api(roles=MANAGER)
def devices_list(request):
    return JsonResponse(_devices_payload(request.user))


@require_POST
@api(roles=MANAGER)
def device_toggle(request, pk):
    dev = Device.objects.select_related("user").get(pk=_uuid(pk))
    if dev.is_active:
        services.revoke_devices(request.user, Device.objects.filter(pk=dev.pk), reason=f"устройство {_name(dev.user)}")
    else:
        services.allow_device(request.user, dev)
    return JsonResponse(_devices_payload(request.user))


@require_POST
@api(roles=MANAGER)
def devices_revoke_all(request):
    if request.data.get("confirm") is not True:
        raise ValidationError("Нужно подтверждение.")
    services.revoke_all_devices(request.user)
    return JsonResponse(_devices_payload(request.user))


# ---------- отчёты ----------

def _period(request):
    try:
        d_to = date.fromisoformat(request.GET.get("to") or timezone.localdate().isoformat())
        d_from = date.fromisoformat(request.GET.get("from") or d_to.isoformat())
    except ValueError as exc:
        raise ValidationError("Неверная дата периода.") from exc
    if d_from > d_to:
        d_from, d_to = d_to, d_from
    if (d_to - d_from).days > 366:
        raise ValidationError("Период — не больше года.")
    return d_from, d_to


def build_report(d_from, d_to):
    start, end = _day_bounds(d_from, d_to)
    shifts = Shift.objects.filter(opened_at__gte=start, opened_at__lt=end)
    rows, totals = [], defaultdict(int)
    for p in Point.objects.filter(is_active=True):
        sh = shifts.filter(point=p)
        sess = Session.objects.filter(shift__in=sh, status=RecordStatus.OK)
        sales = Sale.objects.filter(shift__in=sh, status=RecordStatus.OK)
        exp = Expense.objects.filter(shift__in=sh, status=RecordStatus.OK)
        staff = _sum(sess, "master_share") + _sum(sess, "dispatcher_share") + _sum(sess, "admin_share") + _sum(sales, "admin_share")
        r = {"point": p.name, "shifts": sh.count(), "sessions_count": sess.count(), "sessions": _sum(sess),
             "sales": _sum(sales), "staff": staff, "expenses": _sum(exp),
             "discrepancy": _sum(sh.filter(status=Shift.Status.CLOSED), "discrepancy"),
             "cash": _sum(sess.filter(payment_method="cash")) + _sum(sales.filter(payment_method="cash")),
             "card": _sum(sess.filter(payment_method="card")) + _sum(sales.filter(payment_method="card")),
             "transfer": _sum(sess.filter(payment_method="transfer")) + _sum(sales.filter(payment_method="transfer"))}
        r["revenue"] = r["sessions"] + r["sales"]
        r["profit"] = calc.profit(revenue=r["revenue"], staff=staff, expenses=r["expenses"])
        rows.append(r)
        for k, v in r.items():
            if k != "point":
                totals[k] += v
    people = defaultdict(lambda: {"accrued": 0, "paid": 0, "sessions": 0})
    for s in Session.objects.filter(shift__in=shifts, status=RecordStatus.OK).select_related("master", "dispatcher"):
        if s.master_id:
            people[("e", s.master_id)]["accrued"] += s.master_share
            people[("e", s.master_id)]["sessions"] += 1
        if s.dispatcher_id:
            people[("e", s.dispatcher_id)]["accrued"] += s.dispatcher_share
            people[("e", s.dispatcher_id)]["sessions"] += 1
    for p in Payout.objects.filter(shift__in=shifts):
        key = ("e", p.employee_id) if p.employee_id else ("a", p.admin_id)
        people[key]["paid"] += p.paid
        if p.admin_id:
            people[key]["accrued"] += p.accrued
    names = {e.id: e for e in Employee.objects.filter(id__in=[k[1] for k in people if k[0] == "e"])}
    admins = {u.id: u for u in User.objects.filter(id__in=[k[1] for k in people if k[0] == "a"])}
    staff_rows = []
    for (kind, pk), v in people.items():
        if kind == "e":
            e = names[pk]
            staff_rows.append({"name": e.name, "role": e.get_kind_display(), **v, "debt_total": services.employee_debt(e)})
        else:
            staff_rows.append({"name": _name(admins[pk]), "role": "Администратор", **v, "debt_total": None})
    staff_rows.sort(key=lambda r: -r["accrued"])
    return {"from": d_from.isoformat(), "to": d_to.isoformat(), "points": rows, "totals": dict(totals),
            "staff": staff_rows}


@require_GET
@api(roles=MANAGER)
def report(request):
    return JsonResponse(build_report(*_period(request)))


@require_GET
@api(roles=MANAGER)
def report_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    rep = build_report(*_period(request))
    wb = Workbook()
    ws = wb.active
    ws.title = "Точки"
    head = ["Точка", "Смен", "Сеансов", "Сеансы, ₽", "Доп. продажи, ₽", "Выручка, ₽", "Наличные, ₽", "Карта, ₽",
            "Перевод, ₽", "Персоналу, ₽", "Расходы, ₽", "Прибыль, ₽", "Расхождения кассы, ₽"]
    keys = ["shifts", "sessions_count", "sessions", "sales", "revenue", "cash", "card", "transfer", "staff",
            "expenses", "profit", "discrepancy"]
    money = {"sessions", "sales", "revenue", "cash", "card", "transfer", "staff", "expenses", "profit", "discrepancy"}
    ws.append([f"Отчёт за {rep['from']} — {rep['to']}"])
    ws.append(head)
    for r in rep["points"] + [dict(rep["totals"], point="Итого")]:
        ws.append([r.get("point")] + [(r.get(k, 0) / 100 if k in money else r.get(k, 0)) for k in keys])
    ws2 = wb.create_sheet("Сотрудники")
    ws2.append(["Имя", "Роль", "Сеансов", "Начислено, ₽", "Выплачено, ₽", "Долг на сегодня, ₽"])
    for r in rep["staff"]:
        ws2.append([r["name"], r["role"], r["sessions"], r["accrued"] / 100, r["paid"] / 100,
                    "" if r["debt_total"] is None else r["debt_total"] / 100])
    for sheet in (ws, ws2):
        for cell in sheet[2 if sheet is ws else 1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 16
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "# ##0.00"
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="report_{rep["from"]}_{rep["to"]}.xlsx"'
    return resp


# ---------- журнал ----------

@require_GET
@api(roles=MANAGER)
def log(request):
    q = (request.GET.get("q") or "").strip()
    qs = AuditLog.objects.select_related("actor")
    if q:
        qs = qs.filter(Q(action__icontains=q) | Q(reason__icontains=q))
    rows = [{"at": e.at.isoformat(), "who": _name(e.actor) if e.actor_id else "Система", "action": e.action,
             "old": e.old, "new": e.new, "reason": e.reason} for e in qs[:300]]
    return JsonResponse({"log": rows})
